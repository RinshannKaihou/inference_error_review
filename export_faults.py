#!/usr/bin/env python3
"""
从 FAULT_MASTER_REFERENCE.zh.md 自动解析全部故障条目，导出为 xlsx + csv。
md 始终是唯一 source of truth；本脚本做单向导出。

用法：
    python3 export_faults.py

输出：
    fault_catalog.xlsx  — 格式化的 Excel（含筛选、冻结首行、列宽优化）
    fault_catalog.csv   — 同数据的 CSV（便于 git diff / 命令行处理）
"""

import re
import csv
import os
from datetime import datetime
from collections import Counter

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from category_map import L2_LOCUS_CATEGORY_MAP, CATEGORY_ORDER
MD_PATH = os.path.join(os.path.dirname(__file__), "FAULT_MASTER_REFERENCE.zh.md")
XLSX_PATH = os.path.join(os.path.dirname(__file__), "fault_catalog.xlsx")
CSV_PATH = os.path.join(os.path.dirname(__file__), "fault_catalog.csv")

# ── Column definitions ──────────────────────────────────────────────
COLUMNS = [
    "category",
    "id",
    "name",
    "coverage",
    "confidence",
    "trigger",
    "trigger_review",
    "trigger_note",
    "source",
    "discovery_date",
    "mechanism",
    "behavior",
]


def parse_metadata_line(line: str) -> dict:
    """Parse the metadata line after each ### heading."""
    meta = {}

    # Coverage
    cov_match = re.search(r"Cov:\s*`([^`]+)`", line)
    if cov_match:
        cov = cov_match.group(1).split("（")[0].split("(")[0].strip()
        meta["coverage"] = cov
    else:
        meta["coverage"] = ""

    # Confidence
    conf_match = re.search(r"置信度\s*`([^`]+)`", line)
    meta["confidence"] = conf_match.group(1).strip() if conf_match else ""

    # Trigger label (vLLM/CUDA). A trailing ⚠ marks the row as pending re-review.
    trig_match = re.search(r"触发\s*`([^`]+)`(⚠)?", line)
    meta["trigger"] = trig_match.group(1).strip() if trig_match else ""
    meta["trigger_review"] = "待复核" if (trig_match and trig_match.group(2)) else ""

    return meta


def extract_field(entry_text: str, field_name: str) -> str:
    """Extract a **bold** field's content from entry text."""
    # NOTE: the bare `\n## ` terminator is required. Without it, an entry that is the
    # last one in a `## §x` section and has no `---` separator lets its final bold field
    # swallow the following section heading (previously leaked into discovery_date for
    # 3.8 / 6.25 / 8.10 / K3.10 / K4.9 / K8.14).
    pattern = rf"- \*\*{re.escape(field_name)}\*\*[：:]?\s*(.*?)(?=\n- \*\*|\n### |\n---\s*\n## |\n## |\Z)"
    match = re.search(pattern, entry_text, re.DOTALL)
    if match:
        return match.group(1).strip()
    return ""


def resolve_category(fault_id: str) -> str:
    """Map fault_id to a unified domain-based category, using only the id prefix."""
    if fault_id.startswith("K"):
        # K-layer entries: extract section number, map via locus
        m = re.match(r"K(\d+)\.", fault_id)
        sec_num = m.group(1) if m else ""
        k_loci = {
            "1": "value-corruption", "2": "index-position",
            "3": "block-slot-alloc", "4": "cross-request",
            "5": "cross-node-transfer", "6": "eviction-compression",
            "7": "arch-specific-cache", "8": "control-flow",
            "9": "GLM+Kimi来源补充", "10": "框架/卸载/共享KV/布局/API缓存",
        }
        locus = k_loci.get(sec_num, "")
        return L2_LOCUS_CATEGORY_MAP.get(locus, "KV缓存/注意力")
    else:
        # Numbered entries (1.1-15.x): extract category number
        cat_num = fault_id.split(".")[0]
        cat_names = {
            "1": "硬件/SDC", "2": "GPU显存", "3": "比特翻转攻击",
            "4": "数值/精度", "5": "分布式/多GPU", "6": "服务栈/缓存",
            "7": "量化", "8": "KV缓存/注意力", "9": "位置编码/RoPE",
            "10": "分词器/模板", "11": "解码/采样", "12": "权重/加载",
            "13": "模型图/架构", "14": "长上下文能力", "15": "幻觉/对齐",
        }
        return cat_names.get(cat_num, "未分类")


def extract_source(source_detail: str) -> str:
    """Extract the primary source reference from the source detail text.
    Priority: SPEC+code (project-internal entries) > URL > arXiv > GitHub issue/PR > paper name.
    """
    # Project-internal entries: SPEC section + code path
    spec = re.search(r"SPEC §\d+\.\d+", source_detail)
    code = re.search(r"src/hsd/\S+?(?=[`\s；，。（])", source_detail)
    if spec or code:
        parts = []
        if spec:
            parts.append(spec.group(0))
        if code:
            parts.append(code.group(0))
        return "；".join(parts)

    # URL
    urls = re.findall(r"https?://[^\s）)，。]+", source_detail)
    if urls:
        return urls[0]

    # arXiv number
    arxiv = re.findall(r"arXiv:(\d{4}\.\d{4,5})", source_detail, re.IGNORECASE)
    if arxiv:
        return f"arXiv:{arxiv[0]}"

    # GitHub issue/PR number
    issues = re.findall(r"#\d+", source_detail)
    if issues:
        return issues[0]

    # Paper title in quotes (only if it looks like a real title, not a comment)
    conf = re.search(r'"([A-Z][^"]{10,80})"', source_detail)
    if conf:
        return conf.group(1)

    # Conference reference (e.g. ICML 2024, NeurIPS 2023, USENIX Security 2025)
    conf2 = re.search(r"(ICML|NeurIPS|USENIX|ACL|EMNLP|ICCV|CVPR|ICLR|HPCA|DSN|FSE)\s+20\d{2}", source_detail)
    if conf2:
        return conf2.group(0)

    return source_detail[:80].rstrip("；，。") if source_detail.strip() else ""


def parse_entries(text: str) -> list[dict]:
    """Parse all ### entries from the markdown text."""
    entry_pattern = re.compile(
        r"^### (\S+) `([^`]+)`[^\n]*\n\n([^\n]+)\n\n(.*?)(?=\n### |\Z)",
        re.MULTILINE | re.DOTALL,
    )

    entries = []
    for match in entry_pattern.finditer(text):
        fault_id = match.group(1)
        name = match.group(2)
        meta_line = match.group(3)
        body = match.group(4)

        # Every ### entry is a first-class catalog entry; none are skipped.

        meta = parse_metadata_line(meta_line)

        # Resolve category directly from fault_id
        category = resolve_category(fault_id)

        # Extract text fields
        mechanism = extract_field(body, "机制")
        source_detail = extract_field(body, "来源")
        behavior = extract_field(body, "行为效应")
        trigger_note = extract_field(body, "触发面")
        discovery = extract_field(body, "发现来源")

        # Extract primary source
        source = extract_source(source_detail)

        entry = {
            "category": category,
            "id": fault_id,
            "name": name,
            "coverage": "variant" if meta.get("coverage", "").startswith("var:") else meta.get("coverage", ""),
            "confidence": meta.get("confidence", ""),
            "trigger": meta.get("trigger", ""),
            "trigger_review": meta.get("trigger_review", ""),
            "trigger_note": trigger_note,
            "migration": meta.get("migration", ""),
            "source": source,
            "discovery_date": discovery,
            "mechanism": mechanism,
            "behavior": behavior,
        }
        entries.append(entry)

    # The entry regex requires the metadata on exactly one line (`([^\n]+)\n\n`). If an
    # edit ever wraps a metadata line, that entry is dropped SILENTLY — no warning, no
    # error, just a short export. Fail loudly instead: an `^### ` heading is a fault
    # entry iff it contains a backtick (296 do, 31 section headings do not).
    expected = len(re.findall(r"^### \S+ `", text, re.MULTILINE))
    if len(entries) != expected:
        parsed = {e["id"] for e in entries}
        all_ids = set(re.findall(r"^### (\S+) `", text, re.MULTILINE))
        raise SystemExit(
            f"FATAL: parsed {len(entries)} entries but found {expected} entry headings.\n"
            f"Dropped ids: {sorted(all_ids - parsed)}\n"
            "Most likely a metadata line wrapped onto a second line."
        )

    return entries


def export_csv(entries: list[dict], path: str):
    """Export to CSV."""
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(entries)
    print(f"CSV exported: {path} ({len(entries)} rows)")


def export_xlsx(entries: list[dict], path: str):
    """Export to formatted Excel with filters, frozen header, column widths."""
    if not HAS_OPENPYXL:
        print("openpyxl not available, skipping xlsx export")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fault Catalog"

    # ── Header styling ──
    header_font = Font(name="Microsoft YaHei", bold=True, size=10, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Write headers
    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # ── Data rows ──
    data_font = Font(name="Microsoft YaHei", size=9)
    data_align = Alignment(vertical="top", wrap_text=True)

    # Category color coding
    cat_fills = {}
    palette = [
        "E9F5FE",  # blue
        "FFF8E1",  # yellow
        "FDE9E9",  # red
        "E8F5E9",  # green
        "F3E5F5",  # purple
        "FFF3E0",  # orange
        "E0F7FA",  # cyan
        "FCE4EC",  # pink
        "F1F8E9",  # light green
        "EDE7F6",  # lavender
        "FFFDE7",  # cream
        "EFEBE9",  # brown
        "E8EAF6",  # indigo
        "E0F2F1",  # teal
        "FBE9E7",  # deep orange
        "F4F4F4",  # grey
    ]
    for i, cat in enumerate(CATEGORY_ORDER):
        cat_fills[cat] = PatternFill(start_color=palette[i % len(palette)],
                                      end_color=palette[i % len(palette)],
                                      fill_type="solid")

    for row_idx, entry in enumerate(entries, 2):
        for col_idx, col_name in enumerate(COLUMNS, 1):
            value = entry.get(col_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font
            cell.alignment = data_align
            cell.border = thin_border

            # Color code category column
            if col_name == "category":
                cat = entry.get("category", "")
                if cat in cat_fills:
                    cell.fill = cat_fills[cat]

    # ── Column widths ──
    col_widths = {
        "category": 16, "id": 8, "name": 35,
        "coverage": 10, "confidence": 11,
        "trigger": 13, "trigger_review": 10, "trigger_note": 50, "migration": 18,
        "source": 35, "discovery_date": 20,
        "mechanism": 60, "behavior": 50,
    }
    for col_idx, col_name in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(col_name, 15)

    # ── Auto-filter + freeze header ──
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(entries) + 1}"
    ws.freeze_panes = "A2"

    # ── Summary sheet ──
    ws2 = wb.create_sheet("Summary")
    summary_font = Font(name="Microsoft YaHei", bold=True, size=11)
    label_font = Font(name="Microsoft YaHei", size=10)
    ws2.cell(row=1, column=1, value="故障编目汇总").font = Font(name="Microsoft YaHei", bold=True, size=14)
    ws2.cell(row=2, column=1, value="导出时间").font = label_font
    ws2.cell(row=2, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M")).font = label_font
    ws2.cell(row=3, column=1, value="总条目").font = label_font
    ws2.cell(row=3, column=2, value=len(entries)).font = label_font

    # Count by category
    row = 5
    ws2.cell(row=row, column=1, value="按类别分布").font = summary_font
    row += 1
    cat_counts = Counter(e["category"] for e in entries)
    for cat in CATEGORY_ORDER:
        if cat in cat_counts:
            ws2.cell(row=row, column=1, value=cat).font = label_font
            ws2.cell(row=row, column=2, value=cat_counts[cat]).font = label_font
            row += 1

    # Count by coverage
    row += 1
    ws2.cell(row=row, column=1, value="按覆盖类型分布").font = summary_font
    row += 1
    cov_counts = Counter(e["coverage"] for e in entries)
    for cov, count in cov_counts.most_common():
        ws2.cell(row=row, column=1, value=cov).font = label_font
        ws2.cell(row=row, column=2, value=count).font = label_font
        row += 1

    # Count by confidence
    row += 1
    ws2.cell(row=row, column=1, value="按置信度分布").font = summary_font
    row += 1
    conf_counts = Counter(e["confidence"] for e in entries)
    for conf, count in conf_counts.most_common():
        ws2.cell(row=row, column=1, value=conf).font = label_font
        ws2.cell(row=row, column=2, value=count).font = label_font
        row += 1

    # Count by trigger label (vLLM/CUDA)
    row += 1
    ws2.cell(row=row, column=1, value="按触发标签分布（vLLM/CUDA）").font = summary_font
    row += 1
    trig_counts = Counter(e["trigger"] for e in entries)
    for trig, count in trig_counts.most_common():
        ws2.cell(row=row, column=1, value=trig).font = label_font
        ws2.cell(row=row, column=2, value=count).font = label_font
        row += 1
    ws2.cell(row=row, column=1, value="其中 ⚠ 待复核").font = label_font
    ws2.cell(row=row, column=2, value=sum(1 for e in entries if e["trigger_review"])).font = label_font
    row += 1

    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 15

    wb.save(path)
    print(f"XLSX exported: {path} ({len(entries)} rows, 2 sheets)")


def main():
    print(f"Reading: {MD_PATH}")
    with open(MD_PATH, "r", encoding="utf-8") as f:
        text = f.read()

    entries = parse_entries(text)
    print(f"Parsed {len(entries)} entries")

    # Validate
    empty_names = [e for e in entries if not e["name"]]
    if empty_names:
        print(f"WARNING: {len(empty_names)} entries have empty name field")

    empty_source = [e for e in entries if not e["source"].strip()]
    if empty_source:
        print(f"NOTE: {len(empty_source)} entries have empty source field")

    # Export
    export_csv(entries, CSV_PATH)
    if HAS_OPENPYXL:
        export_xlsx(entries, XLSX_PATH)
    else:
        print("openpyxl not installed — only CSV exported.")

    print(f"\nDone! {len(entries)} entries exported.")
    print(f"  XLSX: {XLSX_PATH}")
    print(f"  CSV:  {CSV_PATH}")


if __name__ == "__main__":
    main()
