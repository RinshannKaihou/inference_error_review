#!/usr/bin/env python3
"""核验推理编目汇总与导出；只读，失败时非零退出。"""
import csv
import re
from collections import Counter
from pathlib import Path

import openpyxl
from export_faults import COLUMNS, parse_entries

ROOT = Path(__file__).resolve().parent


def check_dates(rows):
    for row in rows:
        value = row["discovery_date"]
        if re.search(r"(?:\n|^)(?:\s*---\s*|\s*#+\s+)", value):
            raise ValueError(f"discovery_date 吞入标题/分隔符: {row.get('id', '?')}")


def check(text):
    entries = parse_entries(text)
    total = len(entries)
    match = re.search(r"## 滚动汇总（共 (\d+) 条", text)
    if not match or int(match[1]) != total:
        raise ValueError(f"滚动汇总总数与正文不一致: {match[1] if match else '缺失'} != {total}")
    ids = [row["id"] for row in entries]
    if len(set(ids)) != total:
        raise ValueError("正文故障 ID 重复")

    axes = ("族（Family）", "触发（vLLM/CUDA）", "置信度", "覆盖（Cov）")
    lines = text.splitlines()
    for axis in axes:
        matches = [line for line in lines if line.startswith("|| **" + axis + "** |")]
        if len(matches) != 1:
            raise ValueError(f"滚动汇总缺少或重复 {axis}")
        values = re.findall(r"\*\*(\d+(?:\+\d+⚠)?)\*\*", matches[0])
        count = sum(sum(map(int, re.findall(r"\d+", v))) for v in values)
        if count != total:
            raise ValueError(f"滚动汇总 {axis} 计数 {count} != {total}")
        if axis == "族（Family）":
            families = Counter()
            for metadata in re.findall(r"^### \S+ `[^`]+`[^\n]*\n\n([^\n]+)", text, re.M):
                family = re.match(r"`([^`]+)`", metadata)
                families[family[1] if family else "源文档未给出"] += 1
            if sum(families.values()) != total:
                raise ValueError("家族解析的条目数与正文不一致")
            for key in ("infra_off", "model_on", "ambiguous", "n/a"):
                got = re.search(r"`" + key + r"` \*\*(\d+)\*\*", matches[0])
                if not got or int(got[1]) != families[key]:
                    raise ValueError(f"家族 {key}: 汇总与正文不符")
        if axis == "置信度":
            counts = Counter(row["confidence"] or "源文档未给出" for row in entries)
            for key in ("documented", "verified", "speculative"):
                got = re.search(r"`" + key + r"` \*\*(\d+)\*\*", matches[0])
                if not got or int(got[1]) != counts[key]:
                    raise ValueError(f"置信度 {key}: 汇总与正文不符")
        if axis == "覆盖（Cov）":
            counts = Counter(row["coverage"] for row in entries)
            for label, field in (("NEW", "NEW"), ("var:x", "variant")):
                got = re.search(r"`" + label + r"` \*\*(\d+)\*\*", matches[0])
                if not got or int(got[1]) != counts[field]:
                    raise ValueError(f"覆盖 {label}: 汇总与正文不符")
        if axis == "触发（vLLM/CUDA）":
            counts = Counter(row["trigger"] for row in entries)
            for key in ("config", "other_stack", "partial", "hw_specific", "yes", "no", "n/a"):
                got = re.search(r"`" + key + r"` \*\*(\d+(?:\+\d+⚠)?)\*\*", matches[0])
                if not got or sum(map(int, re.findall(r"\d+", got[1]))) != counts[key]:
                    raise ValueError(f"触发 {key}: 汇总与正文不符")
    directory = text.split("## 目录", 1)[1].split("## 图例", 1)[0]
    numbered = dict((key, int(value)) for key, value in re.findall(
        r"^\|\| (\d{1,2}) \|[^\n]*? \| (\d+) \|$", directory, re.M))
    kv = dict((key, int(value)) for key, value in re.findall(
        r"^\|+ K§(\d+) \|[^\n]*? \| (\d+) \|$", directory, re.M))
    actual_numbered = Counter(row["id"].split(".", 1)[0] for row in entries if not row["id"].startswith("K"))
    actual_kv = Counter(row["id"][1:].split(".", 1)[0] for row in entries if row["id"].startswith("K"))
    if len(numbered) != 15 or numbered != dict(actual_numbered) or len(kv) != 10 or kv != dict(actual_kv):
        raise ValueError("目录分类统计与正文不一致")
    check_dates(entries)
    return entries


def main():
    entries = check((ROOT / "FAULT_MASTER_REFERENCE.zh.md").read_text())
    with (ROOT / "fault_catalog.csv").open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != COLUMNS:
            raise ValueError("CSV 列名与导出脚本不一致")
        rows = list(reader)
    if rows != [{key: entry[key] for key in COLUMNS} for entry in entries]:
        raise ValueError("CSV 与主文档条目不一致；先运行 export_faults.py")
    book = openpyxl.load_workbook(ROOT / "fault_catalog.xlsx", read_only=True, data_only=True)
    sheet = book["Fault Catalog"]
    xrows = sheet.iter_rows(values_only=True)
    if list(next(xrows)) != COLUMNS:
        raise ValueError("XLSX 列名与导出脚本不一致")
    for idx, (actual, expected) in enumerate(zip(xrows, rows), 1):
        # openpyxl reads an empty string cell back as None.
        if [value if value is not None else "" for value in actual] != [expected[k] for k in COLUMNS]:
            raise ValueError(f"XLSX 第 {idx} 行与 CSV 不一致")
    if sheet.max_row != len(rows) + 1:
        raise ValueError("XLSX/CSV 行数不一致")
    book.close()
    print(f"PASS {len(rows)} 条：正文汇总、CSV、XLSX、字段边界一致")


if __name__ == "__main__":
    main()
